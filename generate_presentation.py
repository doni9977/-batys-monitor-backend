from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# We will use a dark blue background if possible, or just standard slides.
# A simple way to make all slides have dark blue background is to define a helper function.

def add_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 30, 60) # Dark Blue
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 200, 50) # Gold
    title.text_frame.paragraphs[0].font.name = 'Arial'
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.placeholders[1]
    content.text = content_text
    for p in content.text_frame.paragraphs:
        p.font.color.rgb = RGBColor(255, 255, 255) # White
        p.font.name = 'Arial'
        p.font.size = Pt(20)
        
    return slide

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(10, 30, 60) # Dark Blue
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255) # White
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle = slide.placeholders[1]
    subtitle.text = subtitle_text
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 200, 50) # Gold
    return slide


add_title_slide(prs, "BATYS MONITOR\nИНТЕЛЛЕКТУАЛЬНАЯ СИСТЕМА МОНИТОРИНГА", "Департамент Экономических Расследований ЗКО")

add_slide(prs, "Архитектура и 3 направления (Домены)", 
          "Единая платформа для контроля сфер экономики:\n\n"
          "• ФСМС / ОСМС: Контроль поликлиник и амбулаторных услуг.\n"
          "• КГД — Нерезиденты: Выявление фиктивных юридических лиц и транзитных номиналов.\n"
          "• Стационар: Контроль больниц и круглосуточных стационаров.\n\n"
          "Мгновенное переключение между модулями в один клик.")

add_slide(prs, "Реестр субъектов и система светофора",
          "Автоматическое ранжирование всех субъектов по уровню риска:\n\n"
          "• Красная зона (Высокий риск): Первоочередные цели для проверок (например, 4+ нарушений).\n"
          "• Желтая зона (Средний риск): Потенциальные подозрения (2-3 нарушения).\n"
          "• Зеленая зона (Низкий риск): Добросовестные субъекты.\n\n"
          "Возможность быстрого экспорта данных в Excel/CSV для дальнейшей работы.")

add_slide(prs, "Загрузка и обработка больших данных",
          "Автоматизированный парсинг сырых отчетов (Excel) из ФСМС и КГД:\n\n"
          "• Не требуется ручной ввод или программирование.\n"
          "• Система сама очищает данные, исправляет форматы и загружает в реляционную базу.\n"
          "• Исключение человеческого фактора на этапе подготовки многомиллионных массивов.")

add_slide(prs, "ИИ-Аналитика (ОСМС и Стационары)",
          "12 уникальных алгоритмов для медицины:\n\n"
          "• S1 (Кросс-чек): Физическая невозможность — пациент одновременно лежит в больнице и посещает поликлинику.\n"
          "• S5 (Мертвые души): Списание услуг и рецептов на имена умерших пациентов.\n"
          "• S2 (Дробление): Искусственное разделение госпитализаций ради двойной оплаты.\n"
          "• A3/A4: Аномальная нагрузка врачей (более 200 пациентов в день).")

add_slide(prs, "ИИ-Аналитика (КГД - Нерезиденты)",
          "Выявление транзитных схем и вывода капитала:\n\n"
          "• Синхронизация с базами ПС КНБ РК.\n"
          "• NR1/NR2: Регистрация компании лицом, не въезжавшим в Казахстан, или въехавшим на 1 день.\n"
          "• NR3 (Сети): Одни и те же нотариусы и переводчики на десятки фиктивных фирм.\n"
          "• NR5: Незаконный вывод валютных ценностей без фактического импорта.")

add_slide(prs, "Автогенерация справок для ДЭР",
          "Формирование готовых рапортов в 1 клик:\n\n"
          "• Документы генерируются в формате Word (.docx).\n"
          "• Использование профессиональной юридической лексики и фабулы.\n"
          "• Детализированные таблицы: ФИО врачей, ИИН пациентов, БИН компаний и суммы ущерба.\n"
          "• Экономия недель рутинной работы аналитиков.")

add_slide(prs, "Аналитика и визуализация (Дашборды)",
          "Интерактивные дашборды для руководства:\n\n"
          "• Визуализация распределения нарушений на графиках.\n"
          "• Быстрый обзор 'ТОП-10' нарушителей по сумме ущерба.\n"
          "• Срезы по регионам, клиникам и периодам для принятия управленческих решений.")

add_slide(prs, "Заключение",
          "Эффективность и перспективы внедрения BatysMonitor:\n\n"
          "• 100% покрытие массива данных (десятки миллионов записей).\n"
          "• Снижение времени на аналитику с месяцев до нескольких минут.\n"
          "• Превентивная блокировка незаконного вывода средств ОСМС и налогов.\n\n"
          "Переход от ручной работы к работе с готовыми инсайтами.")

prs.save('../BatysMonitor_Презентация.pptx')
print("Presentation saved!")
