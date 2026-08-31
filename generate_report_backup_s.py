import sys
import json
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def get_indicator_title(indicator):
    titles = {
        "A1": "Возрастные ограничения",
        "A2": "Гендерные ограничения",
        "A3": "Сверхнагрузка врача",
        "A4": "Превышение дневного лимита",
        "A7": "Превышение годового лимита",
        "A8": "Завышение стоимости",
        "A10": "Нарушение интервала",
        "NR1": "Фиктивное присутствие",
        "NR2": "Транзитный туризм",
        "NR3": "Аффилированные сети",
        "NR4": "Финансовая пустышка",
        "NR5": "Финансовая неактивность",
        "S1": "Кросс-чек поликлиника/стационар",
        "S2": "Дробление госпитализации",
        "S3": "Фиктивный стационар",
        "S4": "Аномальная экстренность",
        "S5": "Услуги после смерти",
    }
    return titles.get(indicator, f"Алгоритм {indicator}")

def get_ai_decision(indicator, d):
    # This is still used for 'S' algorithms
    try:
        if indicator == "S1":
            return f"Услуга '{d.get('service_name', '—')}' оказана в поликлинике {d.get('service_date', '')}, когда пациент находился в стационаре ({d.get('admission_date', '')} - {d.get('discharge_date', '')}). Физически невозможно."
        elif indicator == "S2":
            return f"Повторная госпитализация через {d.get('gap_days', '')} дн. с тем же диагнозом ({d.get('icd10_code', '')}). Предыдущая выписка: {d.get('prev_discharge', '')}, новое поступление: {d.get('new_admission_date', '')}. Признак дробления случая."
        elif indicator == "S3":
            return f"Круглосуточный стационар при пребывании {d.get('bed_days', '')} койко-дн. — дорогой тариф не соответствует сроку. Диагноз: {d.get('diagnosis', '—')} ({d.get('icd10_code', '')})."
        elif indicator == "S4":
            return f"{d.get('reason', '')}. Отделение: {d.get('department', '')}. Экстренных {d.get('emergency_patients', '')} из {d.get('total_patients', '')} ({d.get('emergency_percent', '')}%)."
        elif indicator == "S5":
            return f"Услуга '{d.get('service_name', '—')}' оказана в поликлинике {d.get('service_date', '')}, после зафиксированной даты смерти пациента ({d.get('death_date', '')})."
    except:
        pass
    return str(d)

def generate_algorithm_report(data, output_path):
    doc = docx.Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    indicator = data.get('indicator', 'Unknown')
    title = get_indicator_title(indicator)
    
    h = doc.add_heading(f'Справка по нарушениям алгоритма {indicator}', 1)
    
    doc.add_paragraph(f'Наименование нарушения: {title}', style='List Bullet')
    doc.add_paragraph(f'Всего выявлено нарушений: {data.get("total_risks", 0)}', style='List Bullet')
    if not indicator.startswith("NR"):
        doc.add_paragraph(f'Общая сумма потенциального ущерба: {data.get("total_amount", 0):,.2f} тенге', style='List Bullet')
    
    risks = data.get('risks', [])
    if not risks:
        doc.add_paragraph('Нарушений по данному алгоритму не найдено.')
        doc.save(output_path)
        return

    p_header = doc.add_paragraph()
    run = p_header.add_run('\nДЕТАЛЬНАЯ ИНФОРМАЦИЯ ПО ВЫЯВЛЕННЫМ РИСКАМ:')
    run.bold = True
    
    # ─── Специальная обработка для ОСМС (A) ───
    if indicator.startswith("A"):
        intro_p = doc.add_paragraph()
        if indicator == "A1":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты оказания услуг, не соответствующих возрастной категории пациентов (согласно таблице ниже). К примеру, взрослым пациентам или новорожденным оказывались услуги, предназначенные для совершенно других возрастных групп.\n\nДанные факты указывают на возможные приписки с целью необоснованного получения выплат из фонда ОСМС.")
        elif indicator == "A2":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты оказания услуг, не соответствующих половой принадлежности пациентов (согласно таблице ниже). Пациентам была оказана специфическая медицинская услуга, предназначенная исключительно для противоположного пола.\n\nДанные факты свидетельствуют о внесении недостоверных сведений в информационные системы и указывают на прямые признаки фиктивного оказания услуг (приписок).")
        elif indicator == "A3":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nПри анализе выявлены факты аномальной нагрузки на медицинский персонал, физически превышающие нормативы рабочего времени. Зафиксированы дни со сверхвысоким количеством принятых пациентов.\n\nФизическая невозможность оказания такого объема медицинских услуг в течение одной рабочей смены свидетельствует о признаках приписок.")
        elif indicator == "A4":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nПри изучении сведений выявлены факты необоснованного дублирования медицинских услуг. Установлено, что пациентам в один и тот же день (либо за короткий промежуток времени) неоднократно выставлялась одна и та же услуга.\n\nУказанное является нарушением стандартов оказания медицинской помощи и указывает на искусственное завышение объема оказанных услуг.")
        elif indicator == "A7":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nВ ходе накопительного анализа оказанных услуг за год выявлены факты превышения физиологически возможных норм на одного пациента. Установлено, что пациентам в течение года оказывались услуги в количестве, превышающем годовой лимит и физические возможности человеческого организма.\n\nДанные отклонения носят систематический характер и свидетельствуют о масштабном формировании фиктивных записей.")
        elif indicator == "A8":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты искусственного завышения стоимости оказанных услуг (upcoding). К примеру, при лечении пациентов применялись коды медицинских услуг с более высоким тарифом, применение которых не было обосновано фактическим возрастом или диагнозом пациента.\n\nДанные действия приводят к нецелевому расходованию бюджетных средств Фонда социального медицинского страхования.")
        elif indicator == "A10":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nВыявлены множественные случаи несоответствия заявленной сложности медицинских услуг их фактической продолжительности. У врачей зафиксировано оказание медицинских услуг с аномально коротким интервалом между пациентами.\n\nНормативная сложность и протокол проведения данных операций подразумевают значительно большую затрату времени. Отчётливо прослеживается строгий интервал между услугами, что указывает на пакетное (массовое) внесение данных в информационные системы задним числом.")

        if indicator in ["A1", "A2"]:
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Table Grid'
            h = table.rows[0].cells
            h[0].text = '№'
            h[1].text = 'Врач'
            h[2].text = 'ИИН пациента'
            h[3].text = 'Возраст / Пол'
            h[4].text = 'Услуга'
            h[5].text = 'Код'
            h[6].text = 'Ущерб (₸)'
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(r.get('doctor_name', ''))
                row[2].text = str(r.get('patient_iin', ''))
                val = str(d.get('patient_age', '')) if indicator == "A1" else str(d.get('patient_gender', ''))
                row[3].text = val
                row[4].text = str(d.get('service_name', ''))
                row[5].text = str(d.get('service_code', ''))
                row[6].text = f"{r.get('amount', 0):,.0f}"

        elif indicator == "A3":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            h = table.rows[0].cells
            h[0].text = '№'
            h[1].text = 'Врач'
            h[2].text = 'Дата'
            h[3].text = 'Оказано услуг'
            h[4].text = 'Лимит'
            h[5].text = 'Ущерб (₸)'
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(r.get('doctor_name', ''))
                row[2].text = str(r.get('risk_date', ''))[:10]
                row[3].text = str(d.get('service_count', d.get('daily_count', '')))
                row[4].text = str(d.get('threshold', 200))
                row[5].text = f"{r.get('amount', 0):,.0f}"

        elif indicator in ["A4", "A7"]:
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Table Grid'
            h = table.rows[0].cells
            h[0].text = '№'
            h[1].text = 'Врач'
            h[2].text = 'ИИН пациента'
            h[3].text = 'Услуга'
            h[4].text = 'Оказано раз'
            h[5].text = 'Лимит'
            h[6].text = 'Ущерб (₸)'
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(r.get('doctor_name', ''))
                row[2].text = str(r.get('patient_iin', ''))
                row[3].text = str(d.get('service_name', ''))
                row[4].text = str(d.get('total_count', d.get('total_quantity', '')))
                row[5].text = str(d.get('allowed_per_day', d.get('allowed_per_year', '')))
                row[6].text = f"{r.get('amount', 0):,.0f}"

        elif indicator == "A8":
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Table Grid'
            h = table.rows[0].cells
            h[0].text = '№'
            h[1].text = 'Врач'
            h[2].text = 'ИИН пациента'
            h[3].text = 'Услуга'
            h[4].text = 'Сумма факт'
            h[5].text = 'Тариф (макс)'
            h[6].text = 'Ущерб (₸)'
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(r.get('doctor_name', ''))
                row[2].text = str(r.get('patient_iin', ''))
                row[3].text = str(d.get('service_name', ''))
                row[4].text = str(d.get('actual_amount', ''))
                row[5].text = str(d.get('allowed_amount', ''))
                row[6].text = f"{r.get('amount', 0):,.0f}"

        elif indicator == "A10":
            table = doc.add_table(rows=1, cols=8)
            table.style = 'Table Grid'
            h = table.rows[0].cells
            h[0].text = '№'
            h[1].text = 'Врач'
            h[2].text = 'ИИН пациента'
            h[3].text = 'Пред. услуга'
            h[4].text = 'Тек. услуга'
            h[5].text = 'Инт. факт'
            h[6].text = 'Инт. норма'
            h[7].text = 'Ущерб (₸)'
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(r.get('doctor_name', ''))
                row[2].text = str(r.get('patient_iin', ''))
                row[3].text = str(d.get('previous_service_code', ''))
                row[4].text = str(d.get('service_code', ''))
                row[5].text = str(d.get('actual_interval_minutes', ''))
                row[6].text = str(d.get('required_interval_minutes', ''))
                row[7].text = f"{r.get('amount', 0):,.0f}"
                
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    # ─── Специальная обработка для Нерезидентов (NR) ───
    elif indicator.startswith("NR"):
        intro_p = doc.add_paragraph()
        if indicator == "NR1":
            intro_p.add_run("Справка по результатам аналитической работы. Установлено, что в указанный период были зарегистрированы нижеперечисленные юридические лица, руководителями и учредителями которых выступают нерезиденты.\n\nВ ходе сопоставления сведений с базами данных ПС КНБ РК установлено, что указанные граждане не пересекали государственную границу Республики Казахстан в период государственной регистрации юридических лиц. Процедура регистрации осуществлялась дистанционно с использованием доверенностей.\n\nВышеуказанные факты свидетельствуют о фиктивном характере создания юридических лиц без намерений осуществлять фактическое руководство компанией.")
        elif indicator == "NR2":
            intro_p.add_run("Справка по результатам аналитической работы. Изучением сведений о пересечении государственной границы установлено, что нерезиденты (согласно таблице ниже) осуществили въезд в РК на одних и тех же транспортных средствах. Срок их пребывания на территории РК составил минимальное количество дней.\n\nВ этот период на их имена были зарегистрированы юридические лица. Дополнительно установлено, что на указанных транспортных средствах границу пересекали и иные нерезиденты с целью массовой регистрации юридических лиц, что указывает на организованный ввоз номинальных руководителей.")
        elif indicator == "NR3":
            intro_p.add_run("Справка по результатам аналитической работы. В ходе проведения анализа выявлена группа аффилированных лиц, оказывающих посреднические услуги по массовой регистрации компаний в интересах нерезидентов.\n\nУстановлено, что одни и те же нотариусы и переводчики (согласно таблице) неоднократно выступали посредниками при регистрации иных рисковых юридических лиц, оформленных на нерезидентов. Данный механизм позволяет нерезидентам создавать юридические лица конвейерным способом, что создает предпосылки для их использования в противоправных схемах.")
        elif indicator == "NR4":
            intro_p.add_run("Справка по результатам аналитической работы. Изучением финансово-хозяйственной деятельности нижеперечисленных компаний, зарегистрированных на нерезидентов, установлены признаки фиктивности.\n\nПо данным информационных систем КГД МФ РК, у товариществ отсутствуют обороты по приобретению и реализации товаров, работ и услуг, налоги не уплачивались (либо уплачены в минимальном размере), количество работников составляет 0 человек.\n\nПри этом, согласно банковским выпискам, обороты по счетам компаний носят аномальный характер. Поступившие средства конвертируются и выводятся за рубеж. Фактический импорт, экспорт и налоговые отчисления отсутствуют, компании фактически являются бездействующими.")
        elif indicator == "NR5":
            intro_p.add_run("Справка по результатам аналитической работы. Установлено, что нижеперечисленные ТОО под руководством нерезидентов осуществляют вывод валютных ценностей за пределы Республики Казахстан.\n\nТовариществами заключены международные контракты на поставку товаров. Компании осуществили перевод денежных средств в адрес нерезидентов на крупные суммы.\n\nПри этом, по данным таможенных систем, фактическая поставка товара на территорию РК не осуществлена, возврат денежных средств не произведен. Вышеуказанные факты указывают на признаки уголовного правонарушения, предусмотренного ст. 235-1 УК РК (Незаконный вывод валютных ценностей).")

        if indicator == "NR1":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№'
            hdr_cells[1].text = 'Дата регистрации'
            hdr_cells[2].text = 'Название ТОО'
            hdr_cells[3].text = 'БИН'
            hdr_cells[4].text = 'ФИО нерезидента'
            hdr_cells[5].text = 'ИИН / Паспорт'
            
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(d.get('reg_date', r.get('risk_date', ''))[:10])
                row[2].text = str(d.get('company_name', r.get('clinic_name', '')))
                row[3].text = str(d.get('bin', r.get('patient_iin', '')))
                row[4].text = str(d.get('director_name', r.get('doctor_name', '')))
                row[5].text = str(d.get('director_iin', 'нет данных'))

        elif indicator == "NR2":
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№'
            hdr_cells[1].text = 'Название ТОО'
            hdr_cells[2].text = 'БИН'
            hdr_cells[3].text = 'ФИО нерезидента'
            hdr_cells[4].text = 'Авто (ГРНЗ)'
            hdr_cells[5].text = 'КПП'
            hdr_cells[6].text = 'Дней в РК'
            
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(d.get('company_name', r.get('clinic_name', '')))
                row[2].text = str(d.get('bin', r.get('patient_iin', '')))
                row[3].text = str(d.get('director_name', r.get('doctor_name', '')))
                row[4].text = str(d.get('vehicle_plate', ''))
                row[5].text = str(d.get('crossing_point', ''))
                row[6].text = str(d.get('stay_days', ''))

        elif indicator == "NR3":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№'
            hdr_cells[1].text = 'Название ТОО'
            hdr_cells[2].text = 'БИН'
            hdr_cells[3].text = 'ФИО нерезидента'
            hdr_cells[4].text = 'Переводчик'
            hdr_cells[5].text = 'Нотариус'
            
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(d.get('company_name', r.get('clinic_name', '')))
                row[2].text = str(d.get('bin', r.get('patient_iin', '')))
                row[3].text = str(d.get('director_name', r.get('doctor_name', '')))
                row[4].text = str(d.get('translator', ''))
                row[5].text = str(d.get('notary', ''))

        elif indicator == "NR4":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№'
            hdr_cells[1].text = 'Название ТОО'
            hdr_cells[2].text = 'БИН'
            hdr_cells[3].text = 'ФИО нерезидента'
            hdr_cells[4].text = 'Уставной капитал'
            hdr_cells[5].text = 'Вид деятельности'
            
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(d.get('company_name', r.get('clinic_name', '')))
                row[2].text = str(d.get('bin', r.get('patient_iin', '')))
                row[3].text = str(d.get('director_name', r.get('doctor_name', '')))
                row[4].text = str(d.get('authorized_capital', ''))
                row[5].text = str(d.get('activity_type', ''))

        elif indicator == "NR5":
            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = '№'
            hdr_cells[1].text = 'Название ТОО'
            hdr_cells[2].text = 'БИН'
            hdr_cells[3].text = 'ФИО нерезидента'
            hdr_cells[4].text = 'Статус счета'
            hdr_cells[5].text = 'Баланс (₸)'
            
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                row[0].text = str(i+1)
                row[1].text = str(d.get('company_name', r.get('clinic_name', '')))
                row[2].text = str(d.get('bin', r.get('patient_iin', '')))
                row[3].text = str(d.get('director_name', r.get('doctor_name', '')))
                row[4].text = str(d.get('account_status', ''))
                row[5].text = str(d.get('balance', ''))
        
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    # ─── Стандартная обработка для Стационара (S) ───
    else:
        for i, r in enumerate(risks):
            details = r.get('details', {})
            service_name = details.get('service_name')
            service_code = details.get('service_code')
            diagnosis = details.get('diagnosis')
            icd10 = details.get('icd10_code')
            
            p = doc.add_paragraph()
            
            org_label = "Организация"
            doc_label = "Врач"
            pat_label = "ИИН пациента"
            
            if indicator.startswith("S"):
                org_label = "Стационар (Больница)"
                doc_label = "Лечащий врач"
                
            p.add_run(f"{i+1}. {org_label}: {r.get('clinic_name', '—')}\n").bold = True
            
            if service_name:
                p.add_run(f"   Услуга: ").bold = True
                p.add_run(f"{service_name}")
                if service_code:
                    p.add_run(f" (Код: {service_code})")
                p.add_run("\n")
            elif diagnosis:
                p.add_run(f"   Диагноз: ").bold = True
                p.add_run(f"{diagnosis}")
                if icd10:
                    p.add_run(f" (Код МКБ-10: {icd10})")
                p.add_run("\n")
            
            if r.get('doctor_name') and r.get('doctor_name') != '—':
                p.add_run(f"   {doc_label}: ").bold = True
                p.add_run(f"{r.get('doctor_name')}\n")
            
            if r.get('patient_iin') and r.get('patient_iin') != '—':
                p.add_run(f"   {pat_label}: ").bold = True
                p.add_run(f"{r.get('patient_iin')}\n")
                
            date_str = r.get('risk_date', '')
            if date_str:
                p.add_run(f"   Дата фиксации нарушения: ").bold = True
                p.add_run(f"{date_str[:10]}\n")
            
            ai_decision = get_ai_decision(indicator, details)
            p.add_run(f"   Решение ИИ (Детали): ").bold = True
            font = p.add_run(f"{ai_decision}\n").font
            font.color.rgb = RGBColor(18, 97, 160)
            
            if r.get('amount', 0) > 0:
                p.add_run(f"   Ущерб: ").bold = True
                run = p.add_run(f"{r.get('amount', 0):,.2f} тенге\n")
                run.bold = True
                run.font.color.rgb = RGBColor(194, 24, 91)

    doc.save(output_path)

def main():
    input_data = sys.stdin.read()
    if not input_data.strip():
        sys.exit(1)
        
    data = json.loads(input_data)
    indicator = data.get('indicator', 'Unknown')
    
    os.makedirs('tmp_reports', exist_ok=True)
    output_path = f"tmp_reports/report_{indicator}.docx"
    
    generate_algorithm_report(data, output_path)
    print(output_path)

if __name__ == "__main__":
    main()
