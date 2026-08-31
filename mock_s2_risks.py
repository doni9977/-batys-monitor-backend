import openpyxl
from datetime import datetime, timedelta

file_path = '../Пролеченные случаи стационара/2026 Талап.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb.active

start_row = 10 + 5 + 5 + 15 # 35
# Currently used 25 rows (10 to 34). So let's start at row 35.

for i in range(5):
    row1 = start_row + (i * 2)
    row2 = start_row + (i * 2) + 1
    
    p_name = f"Дробленов Д.Д. {i}"
    p_dob = "15.05.1980"
    icd10 = "I10" # Hypertension
    
    # First admission: 2023-05-01 to 2023-05-10
    adm1 = datetime(2023, 5, 1)
    dis1 = datetime(2023, 5, 10)
    
    # Second admission: 2023-05-11 (1 day gap) to 2023-05-20
    adm2 = datetime(2023, 5, 11)
    dis2 = datetime(2023, 5, 20)
    
    # Row 1 (First hospitalization)
    ws.cell(row=row1, column=3).value = p_name
    ws.cell(row=row1, column=4).value = p_dob
    ws.cell(row=row1, column=6).value = adm1
    ws.cell(row=row1, column=8).value = dis1
    ws.cell(row=row1, column=13).value = icd10 # Код МКБ-10 (Col M = 13)
    ws.cell(row=row1, column=31).value = "Врач-Дробитель"
    
    # Row 2 (Second hospitalization)
    ws.cell(row=row2, column=3).value = p_name
    ws.cell(row=row2, column=4).value = p_dob
    ws.cell(row=row2, column=6).value = adm2
    ws.cell(row=row2, column=8).value = dis2
    ws.cell(row=row2, column=13).value = icd10
    ws.cell(row=row2, column=31).value = "Врач-Дробитель"

wb.save(file_path)
