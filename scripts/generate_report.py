import sys
import json
import docx
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import datetime


def get_indicator_title(indicator):
    titles = {
        "A1": "Нарушение возрастных ограничений при оказании медицинских услуг",
        "A2": "Нарушение гендерных ограничений при оказании медицинских услуг",
        "A3": "Сверхнормативная нагрузка на медицинский персонал",
        "A4": "Превышение дневного лимита кратности услуг",
        "A7": "Превышение годового лимита кратности услуг",
        "A8": "Завышение стоимости медицинских услуг (Upcoding)",
        "A10": "Нарушение нормативного интервала между услугами",
        "NR1": "Фиктивное присутствие нерезидента при регистрации ТОО",
        "NR2": "Транзитный туризм — массовый заезд нерезидентов",
        "NR3": "Аффилированные посреднические сети",
        "NR4": "Финансовая пустышка — отсутствие хозяйственной деятельности",
        "NR5": "Финансовая неактивность — отсутствие банковских операций",
        "S1": "Кросс-чек: услуга в поликлинике во время госпитализации",
        "S2": "Дробление случая госпитализации",
        "S3": "Фиктивный круглосуточный стационар",
        "S4": "Аномально высокий процент экстренных госпитализаций",
        "S5": "Оказание услуг после зафиксированной даты смерти пациента",
    }
    return titles.get(indicator, f"Алгоритм {indicator}")


def set_cell_shading(cell, color):
    """Устанавливает фоновый цвет ячейки."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, size=9, align=None):
    """Устанавливает текст ячейки с форматированием."""
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'


def make_header_row(table, headers):
    """Форматирует заголовок таблицы."""
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, size=9)
        set_cell_shading(cell, "1F4E79")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)


def fmt_amount(val):
    """Форматирует сумму."""
    try:
        v = float(val)
        if v == 0:
            return "—"
        return f"{v:,.0f}".replace(",", " ")
    except:
        return str(val)


def patient_display(r):
    """Возвращает отображение пациента: ФИО (ИИН) или просто ИИН."""
    d = r.get('details', {})
    name = d.get('patient_name', '')
    iin = r.get('patient_iin', d.get('patient_iin', ''))
    if name and name != '—' and name != iin:
        return f"{name} ({iin})"
    return str(iin) if iin else "—"


def generate_algorithm_report(data, output_path):
    doc = docx.Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(4)

    indicator = data.get('indicator', 'Unknown')
    title = get_indicator_title(indicator)
    risks = data.get('risks', [])
    total_amount = data.get('total_amount', 0)
    today = datetime.now().strftime("%d.%m.%Y")

    # ═══════════════════════════════════════════
    #  ЗАГОЛОВОК ДОКУМЕНТА
    # ═══════════════════════════════════════════
    h = doc.add_heading(f'СПРАВКА', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f'по результатам аналитической работы\nАлгоритм {indicator}: {title}')
    run.font.size = Pt(13)
    run.bold = True

    # Мета-информация
    meta_p = doc.add_paragraph()
    meta_p.add_run(f'Дата формирования: ').bold = True
    meta_p.add_run(f'{today}\n')
    meta_p.add_run(f'Всего выявлено нарушений: ').bold = True
    meta_p.add_run(f'{len(risks)}\n')

    if indicator == "A10":
        total_missed = 0
        for r in risks:
            d = r.get('details', {})
            req = float(d.get('required_interval_minutes', 0))
            act = float(d.get('actual_interval_minutes', 0))
            if req > act:
                total_missed += (req - act)
        meta_p.add_run(f'Суммарное отклонение от норматива: ').bold = True
        meta_p.add_run(f'{int(total_missed):,} минут\n'.replace(',', ' '))
    elif indicator == "A3":
        meta_p.add_run(f'Общее количество услуг сверх нормы: ').bold = True
        meta_p.add_run(f'{int(total_amount):,}\n'.replace(',', ' '))
    elif indicator not in ["A1", "A2"] and not indicator.startswith("NR"):
        meta_p.add_run(f'Общая сумма потенциального ущерба: ').bold = True
        meta_p.add_run(f'{total_amount:,.2f} тенге\n')

    if not risks:
        doc.add_paragraph('Нарушений по данному алгоритму не обнаружено.')
        doc.save(output_path)
        return

    # ═══════════════════════════════════════════
    #  ВСТУПИТЕЛЬНЫЙ ТЕКСТ (СПРАВКА)
    # ═══════════════════════════════════════════
    doc.add_paragraph()  # отступ

    intro_texts = {
        "A1": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты оказания услуг, не соответствующих возрастной категории пациентов (согласно таблице ниже). Взрослым пациентам или новорожденным оказывались услуги, предназначенные для совершенно других возрастных групп.\n\nДанные факты указывают на возможные приписки с целью необоснованного получения выплат из фонда ОСМС.",
        "A2": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты оказания услуг, не соответствующих половой принадлежности пациентов (согласно таблице ниже). Пациентам была оказана специфическая медицинская услуга, предназначенная исключительно для противоположного пола.\n\nДанные факты свидетельствуют о внесении недостоверных сведений в информационные системы и указывают на прямые признаки фиктивного оказания услуг (приписок).",
        "A3": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nПри анализе выявлены факты аномальной нагрузки на медицинский персонал, физически превышающие нормативы рабочего времени. Зафиксированы дни со сверхвысоким количеством принятых пациентов.\n\nФизическая невозможность оказания такого объема медицинских услуг в течение одной рабочей смены свидетельствует о признаках приписок.",
        "A4": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nПри изучении сведений выявлены факты необоснованного дублирования медицинских услуг. Установлено, что пациентам в один и тот же день неоднократно выставлялась одна и та же услуга.\n\nУказанное является нарушением стандартов оказания медицинской помощи и указывает на искусственное завышение объема оказанных услуг.",
        "A7": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nВ ходе накопительного анализа оказанных услуг за год выявлены факты превышения физиологически возможных норм на одного пациента. Установлено, что пациентам в течение года оказывались услуги в количестве, превышающем годовой лимит.\n\nДанные отклонения носят систематический характер и свидетельствуют о масштабном формировании фиктивных записей.",
        "A8": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nУстановлено, что выявлены факты искусственного завышения стоимости оказанных услуг (upcoding). При лечении пациентов применялись коды медицинских услуг с более высоким тарифом, применение которых не было обосновано фактическим возрастом или диагнозом пациента.\n\nДанные действия приводят к нецелевому расходованию бюджетных средств Фонда социального медицинского страхования.",
        "A10": "В ходе мониторинга оказания медицинских услуг в рамках ОСМС изучены сведения по медицинским организациям.\n\nВыявлены множественные случаи несоответствия заявленной сложности медицинских услуг их фактической продолжительности. У врачей зафиксировано оказание медицинских услуг с аномально коротким интервалом между пациентами.\n\nНормативная сложность и протокол проведения данных операций подразумевают значительно большую затрату времени. Отчётливо прослеживается строгий интервал между услугами, что указывает на пакетное (массовое) внесение данных в информационные системы задним числом.",
        "NR1": "Установлено, что в указанный период были зарегистрированы нижеперечисленные юридические лица, руководителями и учредителями которых выступают нерезиденты.\n\nВ ходе сопоставления сведений с базами данных ПС КНБ РК установлено, что указанные граждане не пересекали государственную границу Республики Казахстан в период государственной регистрации юридических лиц. Процедура регистрации осуществлялась дистанционно с использованием доверенностей.\n\nВышеуказанные факты свидетельствуют о фиктивном характере создания юридических лиц без намерений осуществлять фактическое руководство компанией.",
        "NR2": "Изучением сведений о пересечении государственной границы установлено, что нерезиденты (согласно таблице ниже) осуществили въезд в РК на одних и тех же транспортных средствах. Срок их пребывания на территории РК составил минимальное количество дней.\n\nВ этот период на их имена были зарегистрированы юридические лица. Дополнительно установлено, что на указанных транспортных средствах границу пересекали и иные нерезиденты с целью массовой регистрации юридических лиц, что указывает на организованный ввоз номинальных руководителей.",
        "NR3": "В ходе проведения анализа выявлена группа аффилированных лиц, оказывающих посреднические услуги по массовой регистрации компаний в интересах нерезидентов.\n\nУстановлено, что одни и те же нотариусы и переводчики (согласно таблице) неоднократно выступали посредниками при регистрации иных рисковых юридических лиц, оформленных на нерезидентов. Данный механизм позволяет нерезидентам создавать юридические лица конвейерным способом, что создает предпосылки для их использования в противоправных схемах.",
        "NR4": "Изучением финансово-хозяйственной деятельности нижеперечисленных компаний, зарегистрированных на нерезидентов, установлены признаки фиктивности.\n\nПо данным информационных систем КГД МФ РК, у товариществ отсутствуют обороты по приобретению и реализации товаров, работ и услуг, налоги не уплачивались (либо уплачены в минимальном размере), количество работников составляет 0 человек.\n\nПри этом, согласно банковским выпискам, обороты по счетам компаний носят аномальный характер. Поступившие средства конвертируются и выводятся за рубеж.",
        "NR5": "Изучением финансово-хозяйственной деятельности компаний, зарегистрированных на нерезидентов, установлены признаки фиктивности и финансовой неактивности.\n\nПо данным банков второго уровня (БВУ), у нижеперечисленных юридических лиц отсутствуют открытые банковские счета, либо на имеющихся счетах наблюдается нулевой (минимальный) остаток средств при отсутствии движения.\n\nОтсутствие банковских счетов или операций по ним свидетельствует о невозможности ведения реальной предпринимательской деятельности.",
        "S1": "В ходе кросс-анализа данных амбулаторной и стационарной помощи установлены факты оказания услуг в поликлиниках пациентам, которые в это время находились на стационарном лечении.\n\nПолучение амбулаторных услуг в период госпитализации физически невозможно, что указывает на фиктивность таких записей.",
        "S2": "При анализе госпитализаций выявлены случаи повторного поступления пациентов с одним и тем же диагнозом в короткий срок после выписки.\n\nДанные факты свидетельствуют о вероятном «дроблении случая» — разделении одной госпитализации на несколько с целью получения дополнительных выплат.",
        "S3": "Выявлены случаи круглосуточной госпитализации с минимальным сроком пребывания, при которых стационарный тариф значительно превышает обоснованный.\n\nКороткий срок пребывания не соответствует тарифу круглосуточного стационара, что указывает на необоснованное завышение расходов.",
        "S4": "Установлены отделения и врачи с аномально высоким процентом экстренных госпитализаций.\n\nПревышение порога экстренности может свидетельствовать о злоупотреблении статусом «экстренная помощь» для обхода плановых квот и увеличения объёма оказанных услуг.",
        "S5": "Выявлены факты оказания медицинских услуг пациентам после зафиксированной даты их смерти.\n\nДанные факты однозначно указывают на фиктивность записей и приписку услуг на умерших граждан.",
    }

    intro = intro_texts.get(indicator, "")
    if intro:
        intro_p = doc.add_paragraph()
        intro_p.add_run("Справка по результатам аналитической работы. ").bold = True
        intro_p.add_run(intro)

    method_texts = {
        "A1": "Сопоставлялись возраст пациента на дату оказания услуги, возрастные границы услуги в классификаторе, код услуги и дата ее оказания. Риск фиксировался, если возраст пациента выходил за допустимый диапазон.",
        "A2": "Сопоставлялись пол пациента, половые ограничения медицинской услуги, код услуги и дата оказания. Риск фиксировался при несоответствии пола пациента установленным ограничениям услуги.",
        "A3": "Сопоставлялись врач, дата оказания, количество услуг за час и за рабочий день с установленными порогами нагрузки. Риск фиксировался, если фактическое количество услуг превышало часовой или дневной норматив.",
        "A4": "Для каждого пациента, врача, услуги и календарного дня подсчитывалось количество оказаний. Полученное значение сравнивалось с допустимой кратностью услуги в день; превышение формировало риск.",
        "A7": "Для каждого пациента и услуги подсчитывалось количество оказаний за календарный год. Это количество сравнивалось с годовым лимитом из классификатора; превышение лимита формировало риск.",
        "A8": "Сравнивались фактическая сумма предъявления, код и количество услуги с максимально допустимым тарифом. Риск фиксировался, если фактическая стоимость превышала расчетную допустимую сумму; разница отражалась как потенциальный ущерб.",
        "A10": "Сопоставлялись время предыдущей и текущей услуги одного врача, фактический интервал между ними и нормативный интервал для текущей услуги. Риск фиксировался, если фактический интервал был меньше требуемого.",
        "NR1": "Сопоставлялись дата регистрации юридического лица, сведения о директоре или учредителе и записи о пересечении государственной границы. Риск фиксировался, если нерезидент не находился в Республике Казахстан в период регистрации компании.",
        "NR2": "Сопоставлялись нерезиденты, даты въезда и выезда, государственные регистрационные данные компаний, номер транспортного средства и пункт пропуска. Риск фиксировался при повторяющихся заездах связанных лиц на одном автомобиле за короткий срок.",
        "NR3": "Сопоставлялись нотариусы, переводчики, нерезиденты и перечень зарегистрированных ими юридических лиц. Риск фиксировался, если один посредник неоднократно участвовал в оформлении нескольких рисковых компаний.",
        "NR4": "Сопоставлялись сведения о регистрации компании с налоговой и финансовой активностью: обороты, уплаченные налоги, численность работников, уставный капитал и движение средств. Риск фиксировался при отсутствии реальной хозяйственной деятельности и наличии признаков транзитного движения денег.",
        "NR5": "Сопоставлялись сведения о компании с данными банков второго уровня: наличие счета, остаток и движение денежных средств. Риск фиксировался при отсутствии счета либо при нулевой или минимальной активности, несовместимой с заявленной деятельностью.",
        "S1": "Сопоставлялись дата и ИИН пациента в амбулаторном реестре с периодом госпитализации в стационарном реестре. Риск фиксировался, если поликлиническая услуга оказана между датами поступления и выписки пациента.",
        "S2": "Сопоставлялись выписка пациента, его повторное поступление, интервал между госпитализациями и код диагноза МКБ-10. Риск фиксировался при повторной госпитализации с тем же диагнозом через короткий срок после выписки.",
        "S3": "Сопоставлялись вид стационара, даты поступления и выписки, количество койко-дней, диагноз и примененный тариф. Риск фиксировался, если круглосуточный стационар предъявлялся при минимальной продолжительности пребывания.",
        "S4": "Для отделения и врача сравнивались общее количество госпитализаций и количество экстренных случаев. Рассчитывалась доля экстренных госпитализаций; риск фиксировался при превышении установленного порога.",
        "S5": "Сопоставлялись дата смерти пациента из регистра и дата оказания медицинской услуги по ИИН. Риск фиксировался, если услуга указана после даты смерти пациента.",
    }
    method = method_texts.get(indicator)
    if method:
        method_p = doc.add_paragraph()
        method_p.add_run("Методика выявления риска. ").bold = True
        method_p.add_run(method)

    # ═══════════════════════════════════════════
    #  СВОДНАЯ ТАБЛИЦА
    # ═══════════════════════════════════════════
    doc.add_paragraph()
    table_title = doc.add_paragraph()
    run = table_title.add_run('Таблица выявленных нарушений:')
    run.bold = True
    run.font.size = Pt(12)

    # ─── ОСМС алгоритмы (A) ───
    if indicator.startswith("A"):
        if indicator == "A1":
            headers = ['№', 'Клиника', 'Врач', 'Пациент', 'ИИН', 'Возраст', 'Допуст. возраст', 'Услуга', 'Код', 'Дата']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], d.get('patient_iin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[5], f"{d.get('patient_age', '—')} лет", size=9)
                min_a = d.get('min_age', 0)
                max_a = d.get('max_age', 0)
                age_range = f"{min_a}–{max_a}" if max_a > 0 else f"от {min_a}"
                set_cell_text(row[6], age_range, size=9)
                set_cell_text(row[7], d.get('service_name', ''), size=9)
                set_cell_text(row[8], d.get('service_code', ''), size=9)
                set_cell_text(row[9], d.get('service_date', r.get('risk_date', ''))[:10], size=9)

        elif indicator == "A2":
            headers = ['№', 'Клиника', 'Врач', 'Пациент', 'ИИН', 'Пол', 'Услуга', 'Код', 'Дата']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], d.get('patient_iin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[5], d.get('patient_gender', ''), size=9)
                set_cell_text(row[6], d.get('service_name', ''), size=9)
                set_cell_text(row[7], d.get('service_code', ''), size=9)
                set_cell_text(row[8], d.get('service_date', r.get('risk_date', ''))[:10], size=9)

        elif indicator == "A3":
            headers = ['№', 'Клиника', 'Врач', 'Дата', 'Услуг за час', 'Норма/час', 'Услуг за день', 'Норма/день']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], str(r.get('risk_date', ''))[:10], size=9)
                set_cell_text(row[4], d.get('service_count', ''), size=9)
                set_cell_text(row[5], d.get('threshold', ''), size=9)
                set_cell_text(row[6], d.get('daily_count', ''), size=9)
                set_cell_text(row[7], '200', size=9)

        elif indicator == "A4":
            headers = ['№', 'Клиника', 'Врач', 'Пациент', 'ИИН', 'Услуга', 'Код', 'Оказано/день', 'Лимит/день', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], d.get('patient_iin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[5], d.get('service_name', ''), size=9)
                set_cell_text(row[6], d.get('service_code', ''), size=9)
                set_cell_text(row[7], d.get('total_count', ''), size=9)
                set_cell_text(row[8], d.get('allowed_per_day', ''), size=9)
                set_cell_text(row[9], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "A7":
            headers = ['№', 'Клиника', 'Врач', 'Пациент', 'ИИН', 'Услуга', 'Код', 'Год', 'Оказано/год', 'Лимит/год', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], d.get('patient_iin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[5], d.get('service_name', ''), size=9)
                set_cell_text(row[6], d.get('service_code', ''), size=9)
                set_cell_text(row[7], d.get('year', ''), size=9)
                set_cell_text(row[8], d.get('total_quantity', ''), size=9)
                set_cell_text(row[9], d.get('allowed_per_year', ''), size=9)
                set_cell_text(row[10], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "A8":
            headers = ['№', 'Клиника', 'Врач', 'Пациент', 'ИИН', 'Услуга', 'Код', 'Кол-во', 'Факт (₸)', 'Тариф макс (₸)', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], d.get('patient_iin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[5], d.get('service_name', ''), size=9)
                set_cell_text(row[6], d.get('service_code', ''), size=9)
                set_cell_text(row[7], d.get('quantity', ''), size=9)
                set_cell_text(row[8], fmt_amount(d.get('actual_amount', '')), size=9)
                set_cell_text(row[9], fmt_amount(d.get('allowed_amount', '')), size=9)
                set_cell_text(row[10], fmt_amount(d.get('excess_amount', r.get('amount', 0))), size=9)

        elif indicator == "A10":
            headers = ['№', 'Клиника', 'Врач', 'Пред. услуга', 'Время пред.', 'Тек. услуга', 'Время тек.', 'Факт (мин)', 'Норма (мин)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('previous_service_name', d.get('previous_service_code', '')), size=9)
                prev_dt = str(d.get('previous_service_date', ''))
                set_cell_text(row[4], prev_dt[11:16] if len(prev_dt) > 16 else prev_dt[:10], size=9)
                set_cell_text(row[5], d.get('service_name', d.get('service_code', '')), size=9)
                svc_dt = str(d.get('service_date', ''))
                set_cell_text(row[6], svc_dt[11:16] if len(svc_dt) > 16 else svc_dt[:10], size=9)
                set_cell_text(row[7], d.get('actual_interval_minutes', ''), size=9)
                set_cell_text(row[8], d.get('required_interval_minutes', ''), size=9)

    # ─── Нерезиденты (NR) ───
    elif indicator.startswith("NR"):
        if indicator == "NR1":
            headers = ['№', 'Дата рег.', 'Название ТОО', 'БИН', 'ФИО нерезидента', 'ИИН/Паспорт']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], str(d.get('reg_date', r.get('risk_date', '')))[:10], size=9)
                set_cell_text(row[2], d.get('company_name', r.get('clinic_name', '')), size=9)
                set_cell_text(row[3], d.get('bin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[4], d.get('director_name', r.get('doctor_name', '')), size=9)
                set_cell_text(row[5], d.get('director_iin', 'нет данных'), size=9)

        elif indicator == "NR2":
            headers = ['№', 'Название ТОО', 'БИН', 'ФИО нерезидента', 'Авто (ГРНЗ)', 'КПП', 'Дней в РК']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], d.get('company_name', r.get('clinic_name', '')), size=9)
                set_cell_text(row[2], d.get('bin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[3], d.get('director_name', r.get('doctor_name', '')), size=9)
                set_cell_text(row[4], d.get('vehicle_plate', ''), size=9)
                set_cell_text(row[5], d.get('crossing_point', ''), size=9)
                set_cell_text(row[6], d.get('stay_days', ''), size=9)

        elif indicator == "NR3":
            headers = ['№', 'Название ТОО', 'БИН', 'ФИО нерезидента', 'Нотариус', 'Переводчик', 'Кол-во компаний']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], d.get('company_name', r.get('clinic_name', '')), size=9)
                set_cell_text(row[2], d.get('bin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[3], d.get('director_name', r.get('doctor_name', '')), size=9)
                set_cell_text(row[4], d.get('notary', ''), size=9)
                set_cell_text(row[5], d.get('translator', ''), size=9)
                set_cell_text(row[6], d.get('pair_count', d.get('companies_count', '')), size=9)

        elif indicator == "NR4":
            headers = ['№', 'Название ТОО', 'БИН', 'ФИО нерезидента', 'Налоги (₸)', 'Сотрудников', 'Устав. капитал (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], d.get('company_name', r.get('clinic_name', '')), size=9)
                set_cell_text(row[2], d.get('bin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[3], d.get('director_name', r.get('doctor_name', '')), size=9)
                set_cell_text(row[4], fmt_amount(d.get('taxes_paid', 0)), size=9)
                set_cell_text(row[5], d.get('employees_count', '0'), size=9)
                set_cell_text(row[6], fmt_amount(d.get('authorized_capital', 0)), size=9)

        elif indicator == "NR5":
            headers = ['№', 'Название ТОО', 'БИН', 'ФИО нерезидента', 'Статус счета', 'Баланс (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], d.get('company_name', r.get('clinic_name', '')), size=9)
                set_cell_text(row[2], d.get('bin', r.get('patient_iin', '')), size=9)
                set_cell_text(row[3], d.get('director_name', r.get('doctor_name', '')), size=9)
                set_cell_text(row[4], d.get('account_status', ''), size=9)
                set_cell_text(row[5], fmt_amount(d.get('balance', 0)), size=9)

    # ─── Стационар (S) ───
    elif indicator.startswith("S"):
        if indicator == "S1":
            headers = ['№', 'Стационар', 'Врач', 'Пациент', 'ИИН', 'Услуга (поликл.)', 'Дата услуги', 'Госпитализация', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], r.get('patient_iin', ''), size=9)
                set_cell_text(row[5], d.get('service_name', ''), size=9)
                set_cell_text(row[6], str(d.get('service_date', ''))[:10], size=9)
                set_cell_text(row[7], f"{d.get('admission_date', '')} — {d.get('discharge_date', '')}", size=9)
                set_cell_text(row[8], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "S2":
            headers = ['№', 'Стационар', 'Врач', 'Пациент', 'ИИН', 'МКБ-10', 'Пред. выписка', 'Нов. поступление', 'Дней между', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], r.get('patient_iin', ''), size=9)
                set_cell_text(row[5], d.get('icd10_code', ''), size=9)
                set_cell_text(row[6], str(d.get('prev_discharge', ''))[:10], size=9)
                set_cell_text(row[7], str(d.get('new_admission_date', ''))[:10], size=9)
                set_cell_text(row[8], d.get('gap_days', ''), size=9)
                set_cell_text(row[9], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "S3":
            headers = ['№', 'Стационар', 'Врач', 'Пациент', 'ИИН', 'МКБ-10', 'Диагноз', 'Койко-дни', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], r.get('patient_iin', ''), size=9)
                set_cell_text(row[5], d.get('icd10_code', ''), size=9)
                set_cell_text(row[6], d.get('diagnosis', ''), size=9)
                set_cell_text(row[7], d.get('bed_days', ''), size=9)
                set_cell_text(row[8], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "S4":
            headers = ['№', 'Стационар', 'Отделение', 'Врач', 'Всего пациентов', 'Экстренных', 'Процент (%)', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], d.get('department', ''), size=9)
                set_cell_text(row[3], r.get('doctor_name', ''), size=9)
                set_cell_text(row[4], d.get('total_patients', ''), size=9)
                set_cell_text(row[5], d.get('emergency_patients', ''), size=9)
                set_cell_text(row[6], d.get('emergency_percent', ''), size=9)
                set_cell_text(row[7], fmt_amount(r.get('amount', 0)), size=9)

        elif indicator == "S5":
            headers = ['№', 'Стационар', 'Врач', 'Пациент', 'ИИН', 'Дата смерти', 'Дата услуги', 'Ущерб (₸)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            make_header_row(table, headers)
            for i, r in enumerate(risks):
                d = r.get('details', {})
                row = table.add_row().cells
                set_cell_text(row[0], i + 1, size=9)
                set_cell_text(row[1], r.get('clinic_name', ''), size=9)
                set_cell_text(row[2], r.get('doctor_name', ''), size=9)
                set_cell_text(row[3], d.get('patient_name', ''), size=9)
                set_cell_text(row[4], r.get('patient_iin', ''), size=9)
                set_cell_text(row[5], str(d.get('death_date', ''))[:10], size=9)
                set_cell_text(row[6], str(d.get('service_date', ''))[:10], size=9)
                set_cell_text(row[7], fmt_amount(r.get('amount', 0)), size=9)

    # ═══════════════════════════════════════════
    #  ПОДПИСЬ
    # ═══════════════════════════════════════════
    doc.add_paragraph()
    doc.add_paragraph()
    sign = doc.add_paragraph()
    sign.add_run('Справка сформирована автоматически системой BatysMonitor (ДЭР ЭКО).').italic = True

    doc.save(output_path)


def main():
    input_data = sys.stdin.read()
    if not input_data.strip():
        sys.exit(1)

    data = json.loads(input_data)
    indicator = data.get('indicator', 'Unknown')

    os.makedirs('tmp_reports', exist_ok=True)
    output_path = f"tmp_reports/report_{indicator}.docx"

    generate_algorithm_report(data, output_path)
    print(output_path)

if __name__ == "__main__":
    main()
